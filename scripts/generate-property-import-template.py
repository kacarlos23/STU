from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


COLUMNS = [
    ("microregionCode", "Código da microrregião", "Obrigatório", "MR01", 20),
    ("street", "Logradouro do imóvel", "Obrigatório", "Rua das Flores", 30),
    ("houseNumber", "Número do imóvel", "Obrigatório", "120", 16),
    ("familyNumber", "Número da família", "Opcional; preencher junto com o responsável", "F-001", 18),
    ("familyResponsibleName", "Nome do responsável pela família", "Opcional; preencher junto com o número da família", "Maria da Silva", 34),
    ("postalCode", "CEP", "Opcional", "45990-000", 16),
    ("complement", "Complemento", "Opcional", "Casa B", 24),
    ("longitude", "Longitude em graus decimais", "Obrigatório", "-39.7419", 18),
    ("latitude", "Latitude em graus decimais", "Obrigatório", "-17.5394", 18),
    ("registrationStatus", "Situação cadastral", "Opcional; Active ou Draft", "Active", 22),
    ("situation", "Situação do imóvel", "Opcional; Occupied, Vacant, Abandoned ou Demolished", "Occupied", 20),
]


def create_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Imóveis"
    headers = [column[0] for column in COLUMNS]
    sheet.append(headers)
    sheet.append([None] * len(headers))
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 32

    for index, (name, description, requirement, example, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, index)
        cell.comment = Comment(f"{description}. {requirement}. Exemplo: {example}.", "STU")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="6D4AFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.cell(2, index).number_format = "@"

    table = Table(displayName="ImportacaoImoveis", ref=f"A1:{get_column_letter(len(headers))}2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
    sheet.add_table(table)

    registration = DataValidation(type="list", formula1='"Active,Draft"', allow_blank=True)
    situation = DataValidation(type="list", formula1='"Occupied,Vacant,Abandoned,Demolished"', allow_blank=True)
    sheet.add_data_validation(registration)
    sheet.add_data_validation(situation)
    registration.add("J2:J10001")
    situation.add("K2:K10001")

    instructions = workbook.create_sheet("Instruções")
    instructions["A1"] = "Como preencher a planilha"
    instructions["A1"].font = Font(bold=True, size=16, color="261F32")
    instructions["A3"] = "Preencha uma linha por imóvel na aba Imóveis e não altere os nomes das colunas."
    instructions["A4"] = "Número e responsável da família devem ser preenchidos juntos ou permanecer ambos vazios."
    instructions["A5"] = "Use ponto ou vírgula como separador decimal para longitude e latitude."
    instructions.append([])
    instructions.append(["Coluna", "Descrição", "Preenchimento", "Exemplo"])
    for name, description, requirement, example, _ in COLUMNS:
        instructions.append([name, description, requirement, example])
    instructions.freeze_panes = "A8"
    instructions_table = Table(displayName="InstrucoesImportacao", ref=f"A7:D{7 + len(COLUMNS)}")
    instructions_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
    instructions.add_table(instructions_table)
    for column, width in zip("ABCD", (24, 42, 48, 24), strict=True):
        instructions.column_dimensions[column].width = width
    for row in instructions.iter_rows(min_row=3, max_row=7 + len(COLUMNS), min_col=1, max_col=4):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)

    validation = load_workbook(path, read_only=False, data_only=False)
    assert validation.sheetnames == ["Imóveis", "Instruções"]
    assert [validation["Imóveis"].cell(1, index).value for index in range(1, len(headers) + 1)] == headers
    validation.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--copy-to", action="append", default=[])
    arguments = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    destinations = [
        root / "src/web/stu-app/public/modelo-importacao-imoveis.xlsx",
        root / "src/web/stu-admin/public/modelo-importacao-imoveis.xlsx",
        *(Path(value).resolve() for value in arguments.copy_to),
    ]
    create_template(destinations[0])
    for destination in destinations[1:]:
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(destinations[0], destination)
    for destination in destinations:
        print(f"{destination} ({destination.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
